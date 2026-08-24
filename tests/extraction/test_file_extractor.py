"""共享文件读取模块测试。"""

from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pymupdf
import pytest
from docx import Document
from PIL import Image
from pydantic import ValidationError

from domain.common import (
    ExtractionMethod,
    ExtractionStatus,
    FilePurpose,
)
from domain.file_models import NormalizedPoint
from extraction import ExtractionConfig, FileExtractor, UnsupportedFileTypeError
from extraction.ocr_reader import RapidOcrReader
from offline.extraction import extract_case_file
from query.extraction import extract_query_file


class FakeOcrEngine:
    """返回固定 RapidOCR v3 风格结果的假引擎。"""

    def __init__(self) -> None:
        """初始化调用次数。"""

        self.call_count = 0

    def __call__(
        self,
        image: np.ndarray,
        *,
        return_word_box: bool,
    ) -> object:
        """返回行级和字级坐标结果。"""

        self.call_count += 1
        height, width = image.shape[:2]
        return SimpleNamespace(
            boxes=np.array(
                [
                    [
                        [0.0, 0.0],
                        [float(width), 0.0],
                        [float(width), float(height) / 2],
                        [0.0, float(height) / 2],
                    ]
                ],
                dtype=np.float32,
            ),
            txts=("风险提示",),
            scores=(0.96,),
            word_results=(
                (
                    (
                        "风险",
                        0.98,
                        [
                            [0.0, 0.0],
                            [float(width) / 2, 0.0],
                            [float(width) / 2, float(height) / 2],
                            [0.0, float(height) / 2],
                        ],
                    ),
                    (
                        "提示",
                        0.94,
                        [
                            [float(width) / 2, 0.0],
                            [float(width), 0.0],
                            [float(width), float(height) / 2],
                            [float(width) / 2, float(height) / 2],
                        ],
                    ),
                ),
            )
            if return_word_box
            else None,
        )


def test_normalized_point_rejects_out_of_page_coordinate() -> None:
    """归一化坐标必须位于页面范围内。"""

    with pytest.raises(ValidationError):
        NormalizedPoint(x=1.1, y=0.5)


def test_rapidocr_reader_preserves_line_and_word_positions() -> None:
    """RapidOCR 行、字坐标和置信度应写入页面草稿。"""

    engine = FakeOcrEngine()
    reader = RapidOcrReader(ExtractionConfig(), engine)
    page = reader.read(Image.new("RGB", (100, 50), "white"))

    assert page.used_ocr is True
    assert page.average_confidence == pytest.approx(0.96)
    assert page.blocks[0].text == "风险提示"
    assert page.blocks[0].polygon == (
        (0.0, 0.0),
        (1.0, 0.0),
        (1.0, 0.5),
        (0.0, 0.5),
    )
    assert [word.text for word in page.words] == ["风险", "提示"]


def test_text_file_builds_offsets_and_checksum(tmp_path: Path) -> None:
    """纯文本读取应生成稳定哈希以及页内和全文偏移。"""

    path = tmp_path / "query.txt"
    path.write_text("第一行\n第二行", encoding="utf-8")

    result = FileExtractor().extract(path, FilePurpose.QUERY)

    assert result.extraction_status is ExtractionStatus.SUCCEEDED
    assert result.extraction_method is ExtractionMethod.PLAIN_TEXT
    assert result.document_id == result.checksum_sha256
    assert result.full_text == "第一行\n第二行"
    assert result.pages[0].blocks[1].page_start_offset == 4
    assert result.pages[0].blocks[1].document_start_offset == 4


def test_pdf_prefers_usable_native_text(tmp_path: Path) -> None:
    """存在足够原生文本时不得调用 RapidOCR。"""

    path = tmp_path / "native.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "This PDF contains enough native text for direct extraction.",
    )
    document.save(path)
    document.close()

    engine = FakeOcrEngine()
    extractor = FileExtractor(
        ExtractionConfig(min_native_text_chars=10),
        engine,
    )
    result = extractor.extract(path, FilePurpose.QUERY)

    assert engine.call_count == 0
    assert result.ocr_page_count == 0
    assert result.pages[0].extraction_method is ExtractionMethod.PDF_TEXT
    assert result.pages[0].blocks[0].polygon is not None


def test_pdf_falls_back_to_rapidocr_for_blank_page(tmp_path: Path) -> None:
    """没有原生文本层的 PDF 页面应使用 RapidOCR。"""

    path = tmp_path / "scan.pdf"
    document = pymupdf.open()
    document.new_page(width=200, height=100)
    document.save(path)
    document.close()

    engine = FakeOcrEngine()
    extractor = FileExtractor(ocr_engine=engine)
    result = extractor.extract(path, FilePurpose.QUERY)

    assert engine.call_count == 1
    assert result.ocr_page_count == 1
    assert result.full_text == "风险提示"
    assert result.pages[0].words[0].page_start_offset == 0
    assert result.pages[0].words[1].page_start_offset == 2
    assert result.extraction_status is ExtractionStatus.SUCCEEDED


def test_docx_reads_paragraphs_and_tables(tmp_path: Path) -> None:
    """DOCX 段落和表格都应进入读取结果。"""

    path = tmp_path / "case.docx"
    document = Document()
    document.add_paragraph("处罚事实")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "机构"
    table.cell(0, 1).text = "罚款"
    document.save(path)

    result = FileExtractor().extract(path, FilePurpose.CASE)

    assert result.extraction_method is ExtractionMethod.DOCX
    assert "处罚事实" in result.full_text
    assert "机构 | 罚款" in result.full_text
    assert result.tables[0].rows == [["机构", "罚款"]]


def test_query_entry_sets_query_purpose_and_logical_uri(tmp_path: Path) -> None:
    """Query 薄入口应设置用途和稳定逻辑路径。"""

    path = tmp_path / "sample.txt"
    path.write_text("查询内容", encoding="utf-8")

    result = extract_query_file(path)

    assert result.file_purpose is FilePurpose.QUERY
    assert result.source_uri == "data/raw/query/sample.txt"


def test_case_entry_sets_case_purpose_and_logical_uri(tmp_path: Path) -> None:
    """案例薄入口应设置案例用途和稳定逻辑路径。"""

    path = tmp_path / "sample.txt"
    path.write_text("案例内容", encoding="utf-8")

    result = extract_case_file(path)

    assert result.file_purpose is FilePurpose.CASE
    assert result.source_uri == "data/raw/cases/sample.txt"


def test_unsupported_extension_is_rejected(tmp_path: Path) -> None:
    """未知扩展名应返回明确异常。"""

    path = tmp_path / "sample.bin"
    path.write_bytes(b"unknown")

    with pytest.raises(UnsupportedFileTypeError):
        FileExtractor().extract(path, FilePurpose.QUERY)
