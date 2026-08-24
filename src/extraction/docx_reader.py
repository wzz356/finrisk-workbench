"""DOCX 文件读取器。"""

from pathlib import Path

from docx import Document

from domain.common import ExtractionMethod
from domain.file_models import ExtractedTable
from extraction.types import DocumentDraft, ElementDraft, PageDraft


def read_docx(path: Path) -> DocumentDraft:
    """读取 DOCX 段落和表格文本；DOCX 无法提供稳定物理页码。"""

    document = Document(str(path))
    blocks = [
        ElementDraft(
            text=paragraph.text.strip(),
            confidence=None,
            polygon=None,
            extraction_method=ExtractionMethod.DOCX,
        )
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    tables: list[ExtractedTable] = []
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append(ExtractedTable(page_number=None, title=None, rows=rows))
        for row in rows:
            row_text = " | ".join(cell for cell in row if cell)
            if row_text:
                blocks.append(
                    ElementDraft(
                        text=row_text,
                        confidence=None,
                        polygon=None,
                        extraction_method=ExtractionMethod.DOCX,
                    )
                )

    warnings = ["DOCX 不提供稳定物理页码，当前统一记录为第 1 页"]
    return DocumentDraft(
        pages=[
            PageDraft(
                width=None,
                height=None,
                blocks=blocks,
                words=[],
                used_ocr=False,
                average_confidence=None,
                quality_score=1.0 if blocks else 0.0,
                extraction_method=ExtractionMethod.DOCX,
                warnings=warnings,
            )
        ],
        tables=tables,
        warnings=warnings,
    )
